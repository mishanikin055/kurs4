"""Create six concise presentation tables from the final comparison CSV files."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from export_comparison_tables import DISPLAY, digest

ROOT = Path(__file__).resolve().parents[1]


def number(value: float, decimals: int) -> str:
    return f"{value:.{decimals}f}".replace(".", ",")


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.2)
    section.left_margin = section.right_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    for name, size in [("Heading 1", 20), ("Heading 2", 13)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string("17365D")
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(4)
    footer = section.footer.paragraphs[0]
    footer.text = "Таблицы для презентации · 09.10.2026 · Лучшие значения выделены зелёным; ↑ больше, ↓ меньше"
    footer.runs[0].font.size = Pt(9)
    sources = {}
    tables = []

    def read(path: str) -> list[dict[str, str]]:
        sources[path] = digest(ROOT / path)
        with (ROOT / path).open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        assert len(rows) == 4 and all(r["status"] == "complete" for r in rows)
        return rows

    def table(title, rows, columns):
        """Columns: label, source key, multiplier, decimal places, best direction."""
        doc.add_heading(title, 2)
        headers = ["Модель"] + [c[0] for c in columns]
        values = [[float(row[c[1]]) * c[2] for c in columns] for row in rows]
        best = [
            (max if column[4] == "max" else min)(values[i][j] for i in range(4))
            if column[4]
            else None
            for j, column in enumerate(columns)
        ]
        body = [
            [DISPLAY.get(row["model"], row["model"])]
            + [number(value, col[3]) for value, col in zip(nums, columns, strict=True)]
            for row, nums in zip(rows, values, strict=True)
        ]
        result = doc.add_table(rows=1, cols=len(headers))
        result.style, result.autofit = "Table Grid", False
        widths = [7.3] + [19.4 / len(columns)] * len(columns)
        for column, width in zip(result.columns, widths, strict=True):
            column.width = Cm(width)
        all_cells = [result.rows[0].cells]
        for line in body:
            cells = result.add_row().cells
            for cell, content in zip(cells, line, strict=True):
                cell.text = content
            all_cells.append(cells)
        for cell, header in zip(all_cells[0], headers, strict=True):
            cell.text = header
        for i, cells in enumerate(all_cells):
            for j, cell in enumerate(cells):
                cell.width = Cm(widths[j])
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(12 if i == 0 else 13)
                if i == 0 or (j and best[j - 1] == values[i - 1][j - 1]):
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "DCE6F1" if i == 0 else "E2EFDA")
                    cell._tc.get_or_add_tcPr().append(shade)
                    for run in paragraph.runs:
                        run.bold = True
        repeat = OxmlElement("w:tblHeader")
        result.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in result.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        tables.append({"title": title, "headers": headers, "rows": body})

    det = read("reports/comparisons/detection.csv")
    doc.add_heading("1. Сравнение детекторов", 1)
    doc.add_paragraph(
        "COCO: 4500 изображений, 4 модели, 3 прохода каждой. RTX 4050, одно изображение за раз."
    )
    table(
        "Качество обнаружения объектов",
        det,
        [
            ("mAP, 0–1 ↑", "mAP", 1, 4, "max"),
            ("AP50, 0–1 ↑", "AP50", 1, 4, "max"),
            ("F1, 0–1 ↑", "F1", 1, 4, "max"),
        ],
    )
    table(
        "Скорость и память",
        det,
        [
            ("Время p50, мс ↓", "p50_ms", 1, 2, "min"),
            ("Время p95, мс ↓", "p95_ms", 1, 2, "min"),
            ("Занятая VRAM, MiB ↓", "VRAM_allocated_MiB", 1, 0, "min"),
        ],
    )
    doc.add_paragraph(
        "Вывод: RF-DETR Small — лучший по качеству; YOLO26s — самый быстрый и требует меньше памяти видеокарты."
    )
    doc.add_paragraph(
        "mAP оценивает категории и расположение рамок; это не процент правильных снимков. AP50 — качество при менее строгом совпадении рамок. F1 — баланс точности и полноты при score 0,5 / IoU 0,5. p50 — типичное время, p95 — время для 95% измерений."
    )

    whole = read("reports/comparisons/classification-validation50000-v1/classification.csv")
    gt = read("reports/comparisons/classification-gt-crops-v1/classification_crops.csv")
    crops = read("reports/comparisons/classification-detector-crops-v1/classification_crops.csv")
    gt_by_model = {r["model"]: r for r in gt}
    merged = [dict(r, gt_top1=gt_by_model[r["model"]]["top1_coarse"]) for r in crops]
    doc.add_page_break()
    doc.add_heading("2. Сравнение классификаторов", 1)
    doc.add_paragraph(
        "Целые изображения: весь ImageNet, 50 000 снимков, 2 прохода. Вырезки: 500 COCO-сцен, 3 ранее завершённых прохода."
    )
    table(
        "Классификация целого изображения: ImageNet",
        whole,
        [
            ("Top-1, % ↑", "top1", 100, 2, "max"),
            ("Top-5, % ↑", "top5", 100, 2, "max"),
            ("Время p50, мс ↓", "p50_ms", 1, 2, "min"),
        ],
    )
    table(
        "Прикладная проверка: вырезки объектов",
        merged,
        [
            ("Top-1: эталонные рамки, % ↑", "gt_top1", 100, 2, "max"),
            ("Top-1: найденные рамки, % ↑", "top1_coarse", 100, 2, "max"),
            ("Macro-F1: найденные рамки ↑", "macro_f1_observed", 1, 4, "max"),
            ("p50: найденные рамки, мс ↓", "p50_ms", 1, 2, "min"),
        ],
    )
    doc.add_paragraph(
        "Вывод: EfficientNetV2-S — лучший Top-1 на ImageNet; ViT-B/16 — лучший Top-1 на вырезках; ResNet-50 — самый быстрый. По Macro-F1 найденных вырезок лидирует EfficientNetV2-S."
    )
    doc.add_paragraph(
        "Top-1 — правильный первый ответ; Top-5 — правильный ответ среди первых пяти. Эталонных объектов 1883; для точности найденных вырезок — 1054 сопоставленных объекта из 2296 рамок. Словарь покрывает 55,81% объектов без crowd; породы не проверены. Замена категории детектора ухудшает результат у всех моделей: классификация показывается отдельно."
    )

    ann = read("reports/comparisons/annotation-test500-v1/annotation.csv")
    doc.add_page_break()
    doc.add_heading("3. Сравнение аннотаторов", 1)
    doc.add_paragraph(
        "500 COCO-изображений, 5 эталонных описаний на снимок. Каждая модель прошла набор один раз; ответы на английском."
    )
    table(
        "Качество текстовых описаний",
        ann,
        [
            ("CIDEr ↑", "CIDEr", 1, 4, "max"),
            ("CHAIRs, % ↓", "CHAIRs", 100, 2, "min"),
            ("Объектная полнота, % ↑", "object_recall_micro", 100, 2, "max"),
        ],
    )
    table(
        "Скорость, память и длина текста",
        ann,
        [
            ("Время p50, с ↓", "total_ms_p50", 0.001, 3, "min"),
            ("Занятая VRAM, MiB ↓", "vram_allocated_peak_bytes", 1 / 1024**2, 0, "min"),
            ("Средняя длина, слов", "words_mean", 1, 1, None),
        ],
    )
    doc.add_paragraph(
        "Вывод: Florence-2 base-ft — лучший для кратких описаний: выше CIDEr, меньше неподтверждённых объектов и быстрее ответ. Qwen3-VL упоминает больше объектов; BLIP требует меньше VRAM."
    )
    doc.add_paragraph(
        "CIDEr — сходство с эталонным текстом; длинные ответы дополнительно штрафуются за отличие длины. CHAIRs — доля описаний с неподтверждённым объектом из словаря COCO, без проверки всех действий и свойств. Полнота — дополнительная доля названных категорий GT. Длина не имеет однозначно лучшего значения. Человеческая оценка исключена; внешний набор не использован."
    )

    directory = ROOT / "reports/exports"
    directory.mkdir(parents=True, exist_ok=True)
    output = directory / "Таблицы для презентации.docx"
    doc.core_properties.title = "Ключевые сравнительные таблицы для презентации"
    doc.core_properties.author = "Курсовой проект"
    doc.save(output)
    reopened = Document(output)
    assert len(reopened.tables) == 6
    for saved, expected in zip(reopened.tables, tables, strict=True):
        assert [[c.text for c in row.cells] for row in saved.rows] == [
            expected["headers"],
            *expected["rows"],
        ]
    manifest = {
        "script_sha256": digest(Path(__file__)),
        "formatting_helper_sha256": digest(ROOT / "scripts/export_comparison_tables.py"),
        "sources": sources,
        "output": "reports/exports/Таблицы для презентации.docx",
        "output_sha256": digest(output),
        "tables": tables,
        "verification": "All six saved tables match the selected CSV values; best values use unrounded numbers",
        "new_inference": False,
        "new_confidence_intervals": False,
    }
    (directory / "presentation_tables_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print("Presentation DOCX created: six tables, all cells verified.")


if __name__ == "__main__":
    build()

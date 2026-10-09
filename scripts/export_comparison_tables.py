"""Export measured comparison tables to Word; requires python-docx 1.2."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DISPLAY = {
    "resnet50": "ResNet-50",
    "efficientnet_v2_s": "EfficientNetV2-S",
    "convnext_tiny": "ConvNeXt-Tiny",
    "vit_b_16": "ViT-B/16",
    "blip_base": "BLIP base",
    "florence2_base_ft": "Florence-2 base-ft",
    "qwen3_vl_2b": "Qwen3-VL 2B Instruct",
    "smolvlm2_500m": "SmolVLM2 500M",
}
CLASSIFIERS = ["ResNet-50", "EfficientNetV2-S", "ConvNeXt-Tiny", "ViT-B/16"]
QUALITY = [
    ("top1", "Top-1, %", "pct"),
    ("top5", "Top-5, %", "pct"),
    ("macro_f1", "Macro-F1, 0–1", "ratio"),
]
TIMING = [
    ("p50_ms", "Типичное время p50, мс", "ms"),
    ("p95_ms", "Время p95, мс", "ms"),
    ("preprocess_mean_ms", "Подготовка входа, среднее, мс", "ms"),
    ("inference_mean_ms", "Работа модели, среднее, мс", "ms"),
    ("postprocess_mean_ms", "Обработка ответа, среднее, мс", "ms"),
    ("read_mean_ms", "Чтение файла, среднее, мс", "ms"),
    ("write_mean_ms", "Запись результата, среднее, мс", "ms"),
    ("cold_load_mean_ms", "Первоначальная загрузка, среднее, мс", "ms"),
    ("RAM_peak_MiB", "Пик RAM, MiB", "mem"),
    ("VRAM_allocated_peak_MiB", "Пик занятой VRAM, MiB", "mem"),
    ("VRAM_reserved_peak_MiB", "Пик зарезервированной VRAM, MiB", "mem"),
    ("parameters", "Число параметров", "int"),
]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fmt(value: str, unit: str = "ratio") -> str:
    if value in {"", "None", "nan"}:
        return "—"
    if unit == "text":
        return value
    number = float(value)
    if unit == "int":
        return f"{int(number):,}".replace(",", " ")
    if unit == "pct":
        number *= 100
    if unit == "bytes":
        number /= 1024**2
    places = 4 if unit == "ratio" else 3 if unit == "pct" else 2 if unit == "ms" else 1
    return f"{number:.{places}f}".replace(".", ",")


class Export:
    def __init__(self, root: Path) -> None:
        self.root = root
        self.sources: dict[str, str] = {}
        self.tables: list[dict] = []
        self.doc = Document()
        section = self.doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21)
        section.top_margin = section.bottom_margin = Cm(1.5)
        section.left_margin = section.right_margin = Cm(1.5)
        section.header_distance = section.footer_distance = Cm(0.6)
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(10)
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.widow_control = True
        for name in ["Title", "Heading 1", "Heading 2"]:
            self.doc.styles[name].font.color.rgb = RGBColor.from_string("17365D")
        section.header.paragraphs[0].text = "Курсовой проект · Итоговые сравнительные таблицы"
        footer = section.footer.paragraphs[0]
        footer.add_run("Результаты на 09.10.2026 · Страница ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        for paragraph in [section.header.paragraphs[0], footer]:
            for run in paragraph.runs:
                run.font.size = Pt(8)

    def read(self, relative: str) -> list[dict[str, str]]:
        path = self.root / relative
        self.sources[relative] = digest(path)
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))

    def paragraph(self, text: str) -> None:
        self.doc.add_paragraph(text)

    def heading(self, text: str, page: bool = False) -> None:
        if page:
            self.doc.add_page_break()
        self.doc.add_heading(text, 1)

    def table(self, title: str, headers: list[str], rows: list[list[str]], source: str) -> None:
        number = len(self.tables) + 1
        self.doc.add_heading(f"Таблица {number}. {title}", 2)
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = False
        width = 26.7 / len(headers)
        for column in table.columns:
            column.width = Cm(width)
        header = table.rows[0]
        repeat = OxmlElement("w:tblHeader")
        header._tr.get_or_add_trPr().append(repeat)
        for cell, text in zip(header.cells, headers, strict=True):
            cell.text = text
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCE6F1")
            cell._tc.get_or_add_tcPr().append(shade)
        # Append XML rows directly: large ImageNet tables avoid quadratic cell lookups.
        template = header._tr
        from copy import deepcopy

        for values in rows:
            row = deepcopy(template)
            props = row.find(qn("w:trPr"))
            for marker in list(props):
                props.remove(marker)
            no_split = OxmlElement("w:cantSplit")
            props.append(no_split)
            for tc, value in zip(row.findall(qn("w:tc")), values, strict=True):
                tc_props = tc.find(qn("w:tcPr"))
                shade = tc_props.find(qn("w:shd"))
                if shade is not None:
                    tc_props.remove(shade)
                for child in list(tc):
                    if child.tag != qn("w:tcPr"):
                        tc.remove(child)
                paragraph = OxmlElement("w:p")
                run = OxmlElement("w:r")
                run_props = OxmlElement("w:rPr")
                size = OxmlElement("w:sz")
                size.set(qn("w:val"), "18")
                run_props.append(size)
                run.append(run_props)
                text = OxmlElement("w:t")
                text.text = value
                run.append(text)
                paragraph.append(run)
                tc.append(paragraph)
            table._tbl.append(row)
        for cell in header.cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)
        self.doc.add_paragraph(f"Источник: {source}", style="Caption")
        self.tables.append({"title": title, "source": source, "headers": headers, "rows": rows})

    def metrics(self, title: str, source: str, fields: list[tuple[str, str, str]]) -> None:
        data = self.read(source)
        assert len(data) == 4 and all(r["status"] == "complete" for r in data)
        headers = ["Показатель"] + [DISPLAY.get(r["model"], r["model"]) for r in data]
        rows = [[label] + [fmt(r[key], unit) for r in data] for key, label, unit in fields]
        self.table(title, headers, rows, source)

    def class_matrix(self, title: str, source: str, names: dict[str, str], key: str) -> None:
        data = self.read(source)
        indexed = {(r[key], DISPLAY.get(r["model"], r["model"])): r for r in data}
        ids = sorted({r[key] for r in data}, key=int)
        rows = []
        for category in ids:
            values = [names[category]]
            for model in CLASSIFIERS:
                row = indexed[category, model]
                values.append(
                    f"{row['correct']}/{row['support']}; "
                    f"{fmt(row['accuracy'], 'pct')}%; F1 {fmt(row['f1'])}"
                    if int(row["support"])
                    else "0/0; точность —; F1 " + fmt(row["f1"])
                )
            rows.append(values)
        self.table(title, ["Категория"] + CLASSIFIERS, rows, source)

    def save(self) -> Path:
        directory = self.root / "reports/exports"
        directory.mkdir(parents=True, exist_ok=True)
        output = directory / "Итоговые таблицы сравнительного анализа.docx"
        self.doc.core_properties.title = "Итоговые таблицы сравнительного анализа моделей"
        self.doc.core_properties.subject = "Детекция, классификация, аннотирование изображений"
        self.doc.core_properties.author = "Курсовой проект Аникина Михаила"
        self.doc.save(output)
        # Verify every cell after saving, including all category rows.
        import zipfile
        from xml.etree import ElementTree as ET

        with zipfile.ZipFile(output) as archive:
            assert archive.testzip() is None
            xml = ET.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        saved_tables = xml.findall(".//w:tbl", ns)
        assert len(saved_tables) == len(self.tables)
        for table, expected in zip(saved_tables, self.tables, strict=True):
            actual = [
                ["".join(cell.itertext()) for cell in row.findall("w:tc/w:p/w:r/w:t", ns)]
                for row in table.findall("w:tr", ns)
            ]
            assert actual == [expected["headers"], *expected["rows"]], expected["title"]
        manifest = {
            "script_sha256": digest(Path(__file__)),
            "sources": self.sources,
            "document": str(output.relative_to(self.root)).replace("\\", "/"),
            "document_sha256": digest(output),
            "tables": [
                {
                    "number": i,
                    "title": table["title"],
                    "source": table["source"],
                    "data_rows": len(table["rows"]),
                    "columns": len(table["headers"]),
                }
                for i, table in enumerate(self.tables, 1)
            ],
            "verification": "All saved table cells match export values from existing CSV",
            "new_inference": False,
            "new_confidence_intervals": False,
            "human_evaluation": "excluded_by_user",
        }
        (directory / "tables_manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        return output


def build(root: Path) -> Path:
    export = Export(root)
    export.doc.add_heading("Итоговые таблицы\nсравнительного анализа моделей", 0)
    export.paragraph(
        "Детекция · Классификация · Аннотирование изображений\nРезультаты на 9 октября 2026 года"
    )
    export.paragraph(
        "Документ объединяет итоговые показатели четырёх моделей в каждой задаче, "
        "измерения времени и памяти, проверку вырезок и результаты по категориям. "
        "Все числа взяты из завершённых локальных экспериментов. "
        "Новые прогоны нейросетей и новые доверительные интервалы для этого файла не выполнялись. "
        "Человеческая оценка аннотаций исключена по решению пользователя."
    )
    export.paragraph(
        "Как читать: доли качества обозначены шкалой 0–1 либо знаком %. "
        "p50 — типичное время, p95 — время, в которое укладываются 95% измерений. "
        "MiB — единица объёма памяти (1 048 576 байт). Занятая и зарезервированная "
        "память PyTorch не равны всей памяти видеокарты. Тире означает, что показатель "
        "не применим или не измерялся. Округление используется только для отображения."
    )
    export.table(
        "Объём сравнений",
        ["Режим", "Разных снимков", "Обработка на модель", "Повторы"],
        [
            ["Детекторы", "4500 COCO", "4500 снимков", "3"],
            ["ImageNet, основное сравнение", "50 000", "50 000 снимков", "2"],
            ["Эталонные вырезки", "500 COCO", "1883 объекта", "3"],
            ["Найденные вырезки", "Те же 500 COCO", "2296 вырезок; 1054 для точности", "3"],
            ["Аннотаторы", "500 COCO", "500 снимков", "1"],
        ],
        "Итоги сравнительного анализа моделей.md и перечисленные ниже CSV",
    )
    export.paragraph(
        "Размеры выбраны до итоговой оценки: у детекции 500 снимков из 5000 COCO "
        "отведены для подготовки, оставшиеся 4500 использованы полностью. "
        "50 000 — весь официальный проверочный ImageNet, по 50 снимков на класс. "
        "500 сцен вырезок дают тысячи объектов. 500 снимков аннотаторов — практический "
        "бюджет медленной генерации текста, нижняя граница предусмотренных 500–1000. "
        "Охвачены 79 из 80 категорий, но редкие классы представлены слабо; формальная "
        "достаточность для любых выводов не заявляется. Повторы используют те же снимки."
    )
    det = "reports/comparisons/detection_full.csv"
    export.heading("1. Детекторы: 4500 COCO-изображений", page=True)
    export.paragraph(
        "Четыре модели, три повтора. Качество рамок и классов оценивается по COCO; mAP не является долей правильно обработанных снимков."
    )
    export.metrics(
        "Качество детекции",
        det,
        [
            (k, label, "ratio")
            for k, label in [
                ("mAP", "COCO mAP"),
                ("AP50", "AP50"),
                ("AP75", "AP75"),
                ("AP_small", "AP: маленькие объекты"),
                ("AP_medium", "AP: средние объекты"),
                ("AP_large", "AP: крупные объекты"),
                ("AR1", "AR, максимум 1 рамка"),
                ("AR10", "AR, максимум 10 рамок"),
                ("AR100", "AR, максимум 100 рамок"),
                ("AR_small", "AR: маленькие объекты"),
                ("AR_medium", "AR: средние объекты"),
                ("AR_large", "AR: крупные объекты"),
            ]
        ],
    )
    export.metrics(
        "Детекция при пороге уверенности 0,5 и совпадении рамок IoU 0,5",
        det,
        [
            ("precision", "Точность найденных объектов", "ratio"),
            ("recall", "Полнота найденных объектов", "ratio"),
            ("F1", "F1", "ratio"),
            ("TP", "Правильно найденные объекты", "int"),
            ("FP", "Неправильные находки", "int"),
            ("FN", "Пропущенные объекты", "int"),
        ],
    )
    export.metrics(
        "Скорость и ресурсы детекторов",
        det,
        [
            ("p50_ms", "p50, мс", "ms"),
            ("p95_ms", "p95, мс", "ms"),
            ("images_per_second", "Пропускная способность, снимков/с", "ms"),
            ("preprocess_mean_ms", "Подготовка, среднее, мс", "ms"),
            ("forward_mean_ms", "Работа модели, среднее, мс", "ms"),
            ("postprocess_mean_ms", "Обработка ответа, среднее, мс", "ms"),
            ("read_mean_ms", "Чтение файла, среднее, мс", "ms"),
            ("write_mean_ms", "Запись, среднее, мс", "ms"),
            ("cold_load_ms_mean", "Первоначальная загрузка, среднее, мс", "ms"),
            ("RAM_peak_MiB", "Пик RAM, MiB", "mem"),
            ("VRAM_allocated_MiB", "Пик занятой VRAM, MiB", "mem"),
            ("VRAM_reserved_MiB", "Пик зарезервированной VRAM, MiB", "mem"),
            ("parameters", "Число параметров", "int"),
            ("weight_MiB", "Размер весов, MiB", "mem"),
        ],
    )
    export.paragraph(
        "Выбор по качеству — RF-DETR Small; кандидат при приоритете скорости — YOLO26s. Перенесены девять завершённых проходов, три RF-DETR восстановлены после исправления чтения списка категорий; ранние сбои сохранены."
    )
    whole = "reports/comparisons/classification-validation50000-v1/classification.csv"
    export.heading("2. Классификация целых снимков: ImageNet 50 000", page=True)
    export.paragraph(
        "1000 классов, по 50 изображений каждого. Качество — первый полный проход; время и ресурсы — одинаковые первые два прохода всех четырёх моделей. Ранее выполненный третий ResNet-50 сохранён, но исключён из основной таблицы."
    )
    export.metrics("Качество классификации ImageNet", whole, QUALITY)
    export.metrics(
        "Скорость и ресурсы классификации ImageNet",
        whole,
        TIMING
        + [
            ("images_per_second", "Пропускная способность, снимков/с", "ms"),
            ("weight_MiB", "Размер весов, MiB", "mem"),
        ],
    )
    export.paragraph(
        "Top-1 — правильный первый ответ; Top-5 — правильный ответ среди первых пяти. На этом наборе лучший Top-1 у EfficientNetV2-S. Штатная подготовка входа и разрешения моделей различаются."
    )
    for mode, title in [
        ("gt", "3. Классификация эталонных вырезок"),
        ("detector", "4. Классификация найденных вырезок"),
    ]:
        source = f"reports/comparisons/classification-{mode}-crops-v1/classification_crops.csv"
        export.heading(title, page=True)
        export.paragraph(
            "Те же 500 COCO-сцен; три уже завершённых повтора. Качество — первый проход. Словари ImageNet и COCO согласованы для 56 категорий; покрытие объектов без crowd — 55,81%. Это оценка общих категорий, а не пород и тонких подтипов."
        )
        fields = [
            ("n_crops", "Всего вырезок", "int"),
            ("n_eligible_crops", "Вырезок в расчёте точности", "int"),
            ("top1_coarse", "Top-1 по общей категории, %", "pct"),
            ("top5_coarse", "Top-5 по общей категории, %", "pct"),
            ("macro_f1_observed", "Macro-F1 наблюдаемых категорий", "ratio"),
            ("acceptance_rate", "Доля ответов, прошедших пороги, %", "pct"),
            ("accepted_top1_accuracy", "Top-1 среди прошедших пороги, %", "pct"),
        ]
        export.metrics("Качество ответов на вырезках", source, fields)
        export.metrics(
            "Скорость и ресурсы на вырезках",
            source,
            TIMING + [("crop_mean_ms", "Вырезание объекта, среднее, мс", "ms")],
        )
        if mode == "detector":
            export.metrics(
                "Проверка замены категории детектора",
                source,
                [
                    ("detector_only_matched_accuracy", "Точность категории детектора, %", "pct"),
                    (
                        "hypothetical_gated_matched_accuracy",
                        "Точность при замене после порогов, %",
                        "pct",
                    ),
                    ("helpful_changes", "Исправленные ошибки", "int"),
                    ("harmful_changes", "Внесённые ошибки", "int"),
                    ("missed_supported_gt", "Пропущенные поддерживаемые объекты", "int"),
                    ("unmatched_proposals", "Рамки без сопоставления с GT", "int"),
                    ("matched_unsupported_proposals", "Неподдерживаемые категории", "int"),
                    ("crowd_ignored_proposals", "Рамки crowd, исключённые из оценки", "int"),
                    (
                        "detector_only_supported_recall",
                        "Полнота правильных категорий детектора, %",
                        "pct",
                    ),
                    ("hypothetical_gated_supported_recall", "Полнота после замены, %", "pct"),
                ],
            )
            export.paragraph(
                "Рамки получены RF-DETR Small. Условная точность считается по 1054 сопоставленным поддерживаемым объектам; пропуски учтены отдельно. Замена ухудшила точность у всех моделей. Категория детектора сохраняется, классификатор даёт отдельный ответ. Пороги score 0,5 / margin 0,1 диагностические, надёжный отказ на dev не откалиброван."
            )
        export.paragraph(
            "По Top-1/Top-5 для вырезок выбран ViT-B/16. Для найденных вырезок Macro-F1 лучше у EfficientNetV2-S; это другой критерий."
        )
    ann = "reports/comparisons/annotation-test500-v1/annotation.csv"
    export.heading("5. Аннотаторы: 500 COCO-изображений", page=True)
    export.paragraph(
        "Четыре модели × один проход × 500 изображений, без ошибок; пять эталонных описаний на снимок. FP16, одно изображение за раз, максимум 96 новых токенов. Описания на английском. Человеческая оценка исключена."
    )
    export.metrics(
        "Качество и длина аннотаций",
        ann,
        [
            ("CIDEr", "CIDEr, исходная шкала evaluator", "ratio"),
            ("BLEU_4", "BLEU-4, 0–1", "ratio"),
            ("ROUGE_L", "ROUGE-L, 0–1", "ratio"),
            ("CHAIRs", "Тексты с неподтверждённым объектом, %", "pct"),
            ("CHAIRi", "Неподтверждённые упоминания объектов, %", "pct"),
            ("object_recall_micro", "Дополнительная объектная полнота, %", "pct"),
            ("words_mean", "Средняя длина, слов", "ms"),
            ("output_tokens_mean", "Средняя длина, токенов", "ms"),
            ("token_budget_reached", "Ответов, достигших лимита 96 токенов", "int"),
        ],
    )
    ann_fields = []
    for field, label in [
        ("preprocess", "Подготовка входа"),
        ("generate", "Генерация текста"),
        ("postprocess", "Обработка ответа"),
        ("total", "Общее время без чтения"),
        ("total_with_read", "Общее время с чтением"),
    ]:
        for suffix, description in [("mean", "среднее"), ("p50", "p50"), ("p95", "p95")]:
            ann_fields.append((f"{field}_ms_{suffix}", f"{label}: {description}, мс", "ms"))
    export.metrics("Время создания аннотаций", ann, ann_fields)
    export.metrics(
        "Ресурсы и скорость генерации аннотаторов",
        ann,
        [
            ("cold_load_ms", "Первоначальная загрузка модели, мс", "ms"),
            ("rss_peak_bytes", "Пик RAM/RSS, MiB", "bytes"),
            ("vram_allocated_peak_bytes", "Пик занятой VRAM, MiB", "bytes"),
            ("vram_reserved_peak_bytes", "Пик зарезервированной VRAM, MiB", "bytes"),
            ("tokens_per_second_mean", "Токенов/с: среднее", "ms"),
            ("tokens_per_second_p50", "Токенов/с: p50", "ms"),
            ("tokens_per_second_p95", "Токенов/с: p95", "ms"),
        ],
    )
    export.paragraph(
        "Выбран Florence-2 base-ft для коротких английских описаний. CIDEr штрафует отличие длины от эталонов; почти нулевой балл длинного Qwen не означает отсутствие смысла в каждом ответе. CHAIR проверяет объекты из словаря, но не все действия и свойства. Внешний набор и SPICE не использованы; известное обучение части моделей на COCO ограничивает выводы о новых данных."
    )
    export.heading("Приложение А. Предварительная оценка ImageNet 5000", page=True)
    prelim = "reports/comparisons/classification-evaluation5000-v1/classification.csv"
    export.paragraph(
        "Четыре модели, три повтора, по пять изображений каждого класса. Этот набор входит в основные 50 000 и не является независимым подтверждением итогов."
    )
    export.metrics("Предварительное качество на 5000 снимках", prelim, QUALITY)
    export.metrics(
        "Предварительные скорость и ресурсы",
        prelim,
        TIMING
        + [("images_per_second", "Снимков/с", "ms"), ("weight_MiB", "Размер весов, MiB", "mem")],
    )
    export.heading("Приложение Б. Детекция по отдельным категориям", page=True)
    det_classes = "reports/comparisons/detection_per_class.csv"
    rows = export.read(det_classes)
    for model in dict.fromkeys(r["model"] for r in rows):
        values = [
            [
                r["category_name"],
                fmt(r["AP"]),
                fmt(r["precision"]),
                fmt(r["recall"]),
                fmt(r["F1"]),
                r["TP"],
                r["FP"],
                r["FN"],
            ]
            for r in rows
            if r["model"] == model
        ]
        export.table(
            model + ": все 80 категорий COCO",
            ["Категория", "AP", "Precision", "Recall", "F1", "TP", "FP", "FN"],
            values,
            det_classes,
        )
    export.heading("Приложение В. ImageNet: все 1000 классов", page=True)
    export.paragraph(
        "Каждая ячейка: правильных ответов / число снимков; Top-1 в процентах; F1. Названия классов сохранены на английском, как в исходном словаре. Качество — первый полный проход."
    )
    labels_path = root / "classification/configs/imagenet_labels.json"
    export.sources[str(labels_path.relative_to(root)).replace("\\", "/")] = digest(labels_path)
    labels = json.loads(labels_path.read_text(encoding="utf-8"))["classes"]
    names = {str(r["index"]): f"{r['index']} · {r['name']}" for r in labels}
    export.class_matrix(
        "Качество всех классификаторов по классам",
        "reports/comparisons/classification-validation50000-v1/classification_per_class.csv",
        names,
        "index",
    )
    export.heading("Приложение Г. Вырезки: все 56 поддерживаемых категорий", page=True)
    export.paragraph(
        "Формат ячеек тот же: правильных / число объектов; Top-1, %; F1. При отсутствии объектов точность не определена. Нулевой F1 берётся из сохранённого отчёта, а не считается доказательством качества отсутствующего класса."
    )
    coco_names = {r["category_id"]: r["category_name"] for r in rows}
    for mode, label in [("gt", "Эталонные рамки"), ("detector", "Найденные рамки")]:
        export.class_matrix(
            label,
            f"reports/comparisons/classification-{mode}-crops-v1/per_class.csv",
            coco_names,
            "category_id",
        )
    export.heading("Приложение Д. Покрытие категорий набора аннотаторов", page=True)
    coverage = "reports/comparisons/annotation-test500-v1/dataset_category_coverage.csv"
    data = export.read(coverage)
    export.paragraph(
        "Число сцен показывает, в скольких из 500 снимков есть категория. Количество объектов включает crowd; crowd указан отдельным столбцом. Один снимок может содержать несколько категорий, поэтому суммы сцен не равны 500."
    )
    export.table(
        "Все 80 категорий COCO",
        ["Категория", "Объектов всего", "Из них crowd", "Сцен"],
        [[r["category"], r["instances"], r["crowd_instances"], r["scenes"]] for r in data],
        coverage,
    )
    export.heading("Приложение Е. Исторические интервалы детекции", page=True)
    export.paragraph(
        "Эти значения рассчитаны ранее, до запрета новых доверительных интервалов. Они перенесены без пересчёта. Для классификаторов и аннотаторов интервалы не добавлялись."
    )
    export.metrics(
        "Ранее сохранённые границы 95%-интервала mAP",
        det,
        [
            ("mAP_ci95_low", "Нижняя граница", "ratio"),
            ("mAP_ci95_high", "Верхняя граница", "ratio"),
        ],
    )
    export.heading("Источники и возможность повторить экспорт", page=True)
    export.paragraph(
        "В документ включены итоговые сравнения моделей и таблицы по категориям. Полные матрицы ошибок, ответы на отдельные изображения, цифровые отпечатки весов и незаполненные формы человеческой оценки хранятся в исходных материалах отчётов; это не новые сравнительные таблицы."
    )
    for source in export.sources:
        export.paragraph(source)
    export.paragraph(
        "Понятное описание всех шагов: «Как выполнялся сравнительный анализ.md». Выводы и причины размеров выборок: «Итоги сравнительного анализа моделей.md». Для повторного экспорта используется scripts/export_comparison_tables.py в окружении с python-docx 1.2.0. Таблицы редактируются в Word как обычные таблицы. Источники и контрольная сумма этого DOCX — reports/exports/tables_manifest.json."
    )
    return export.save()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    print(build(args.root.resolve()))


if __name__ == "__main__":
    main()
